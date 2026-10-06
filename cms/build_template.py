"""Build the starter content workbook (upload to Google Drive; it converts to Google Sheets).

    npx tsx cms/export-defaults.ts > /tmp/defaults.json
    python3 cms/build_template.py /tmp/defaults.json cms/Elderflowers-CMS.xlsx
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PLUM, CREAM, BUTTER = "3E2330", "F5F0E6", "EADFB6"
HEAD_FONT = Font(bold=True, color=CREAM)
HEAD_FILL = PatternFill("solid", fgColor=PLUM)
NOTE_FONT = Font(italic=True, color="6E5A62")

GUIDE = [
    ("CÁCH SỬA WEBSITE ELDERFLOWERS GARDEN", ""),
    ("", ""),
    ("1. Sửa chữ, giá, sản phẩm", "Sửa thẳng vào các tab bên dưới. Website tự cập nhật sau khoảng 1–2 phút (tải lại trang 2 lần)."),
    ("2. KHÔNG đổi tên tab hoặc tên cột", "Dòng đầu tiên (màu tím) là tên cột. Đổi tên thì website không đọc được cột đó."),
    ("3. Thêm sản phẩm", "Thêm một dòng mới ở tab SanPham. Cột Mã: viết liền, không dấu, không trùng (VD: hoa-hong-box)."),
    ("4. Ẩn sản phẩm", "Ở cột Hiển thị chọn Không. Không nên xoá dòng nếu sản phẩm có thể bán lại."),
    ("5. Giá", "Gõ số, có hoặc không có dấu chấm đều được: 220000 hoặc 220.000."),
    ("6. Thêm / thay ảnh", "Tải ảnh lên thư mục 'Elderflowers – Ảnh website' trên Google Drive → bấm chuột phải → Chia sẻ → Sao chép đường liên kết → dán vào cột Ảnh (link). Ảnh sẽ tự được công khai để website hiển thị."),
    ("   Kích thước ảnh", "Ảnh nên rộng 1600–2000px, dưới 1MB. Sản phẩm: dọc 4:5. Đầu trang: ngang 5:4. Hộp hoa: vuông. Ghé vườn: ngang 16:10."),
    ("7. Tháng có mùa", "Ghi số tháng cách nhau bằng dấu phẩy, VD: 10, 11, 12, 1, 2. Hoặc ghi: Quanh năm."),
    ("8. Nguyên liệu (mã)", "Ghi Mã của nguyên liệu ở tab NguyenLieu, nhiều cái thì cách nhau dấu phẩy, VD: elder, lemon."),
    ("9. Đơn hàng", "Đơn mới tự xuất hiện ở tab DonHang và được gửi email. Đổi cột Trạng thái khi xử lý (Mới → Đã gọi → Đã giao)."),
    ("10. Cài đặt chung", "Tab CaiDat: banner, tiêu đề, phí giao hàng, tài khoản ngân hàng cho mã QR. Chỉ sửa cột Giá trị."),
    ("", ""),
    ("Bảo mật", "Không chia sẻ bảng tính này công khai: nó chứa tên, số điện thoại và địa chỉ của khách."),
]


def header(ws, headers, widths=None):
    ws.append(headers)
    for i, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(headers[i - 1], 22)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28


def main(src, dest):
    data = json.load(open(src, encoding="utf-8"))
    schema = data["schema"]
    wb = Workbook()

    g = wb.active
    g.title = "HuongDan"
    for row in GUIDE:
        g.append(row)
    g["A1"].font = Font(bold=True, size=14, color=PLUM)
    for r in range(3, len(GUIDE) + 1):
        g.cell(row=r, column=1).font = Font(bold=True)
        g.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    g.column_dimensions["A"].width = 34
    g.column_dimensions["B"].width = 110
    g.sheet_properties.tabColor = BUTTER

    s = wb.create_sheet(schema["settingsSheet"])
    cols = schema["settingsColumns"]
    header(s, [cols["key"], cols["value"], cols["note"]], {cols["key"]: 22, cols["value"]: 70, cols["note"]: 70})
    for key, value, note in data["settings"]:
        s.append([key, value, note])
        s.cell(row=s.max_row, column=1).font = Font(color="6E5A62")
        s.cell(row=s.max_row, column=2).alignment = Alignment(wrap_text=True)
        s.cell(row=s.max_row, column=3).font = NOTE_FONT
    s.sheet_properties.tabColor = BUTTER

    wide = {"Câu chuyện": 70, "Nội dung": 80, "Cách làm": 50, "Mô tả": 45, "Mô tả ngắn": 40, "Ảnh (link)": 45, "Tên (EN)": 28, "Tên (VI)": 28, "Nơi trồng": 32}
    for sheet, tab in data["tabs"].items():
        ws = wb.create_sheet(sheet)
        header(ws, tab["headers"], wide)
        for row in tab["rows"]:
            ws.append(row)
        for r in ws.iter_rows(min_row=2):
            for cell in r:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if sheet == schema["tabs"]["products"]["sheet"]:
            cols = schema["tabs"]["products"]["columns"]
            for field, choices in (("category", data["categories"]), ("visible", ["Có", "Không"])):
                letter = get_column_letter(tab["headers"].index(cols[field]) + 1)
                dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}500")

    for rec in schema["records"].values():
        ws = wb.create_sheet(rec["sheet"])
        header(ws, rec["columns"], {"Sản phẩm": 40, "Địa chỉ": 40, "Lời nhắn": 30})
        ws.sheet_properties.tabColor = "55654A"

    wb.save(dest)
    print("wrote", dest)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
